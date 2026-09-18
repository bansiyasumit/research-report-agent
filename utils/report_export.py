"""
Converts the final markdown report into downloadable PDF and DOCX files
for the Streamlit download buttons.
Filled in Step 8 of the build.
"""


import re

from fpdf import FPDF
from fpdf.enums import XPos, YPos

# fpdf2's built-in core fonts (Helvetica etc.) only support Latin-1, but
# LLM-written text routinely uses smart quotes, em/en dashes, ellipses,
# and bullet characters that fall outside it. Map the common ones to
# ASCII equivalents rather than bundling a Unicode TTF font just for this.
_PDF_CHAR_MAP = {
    "‘": "'", "’": "'", "“": '"', "”": '"',
    "–": "-", "—": "-", "…": "...", "•": "*",
    " ": " ",
}


def _sanitize_for_pdf(text: str) -> str:
    for unicode_char, ascii_char in _PDF_CHAR_MAP.items():
        text = text.replace(unicode_char, ascii_char)
    # Anything else outside Latin-1 (e.g. emoji) is dropped rather than
    # crashing the export.
    return text.encode("latin-1", errors="ignore").decode("latin-1")


def _heading_level(line: str) -> int:
    """Returns 1-6 for a Markdown heading line, 0 if it isn't one."""
    stripped = line.lstrip("#")
    level = len(line) - len(stripped)
    if 1 <= level <= 6 and stripped.startswith(" "):
        return level
    return 0


def markdown_to_pdf(markdown_text: str, output_path: str) -> str:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # Font sizes step down as heading level increases (H1 biggest, H6 smallest),
    # capped at 11pt so deep headings don't shrink below body text.
    heading_sizes = {1: 20, 2: 16, 3: 14, 4: 12, 5: 11, 6: 11}

    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()

        if not line:
            pdf.ln(4)
            continue

        level = _heading_level(line)
        if level:
            # Headings are already bold via set_font, so markdown=True is
            # skipped here (it would toggle bold OFF on hitting "**") —
            # any stray ** in a heading is just stripped instead.
            text = _sanitize_for_pdf(line.lstrip("#").strip()).replace("**", "")
            pdf.set_font("Helvetica", style="B", size=heading_sizes[level])
            pdf.multi_cell(0, 8 + (6 - level), text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(1)
            continue

        stripped = line.strip()
        pdf.set_font("Helvetica", size=11)
        if stripped.startswith(("- ", "* ")):
            # markdown=True renders inline **bold** within the bullet text
            # (e.g. "- **Key finding:** ...") instead of showing literal
            # asterisks.
            pdf.multi_cell(
                0, 6, _sanitize_for_pdf(f"    * {stripped[2:]}"),
                new_x=XPos.LMARGIN, new_y=YPos.NEXT, markdown=True,
            )
            continue

        pdf.multi_cell(
            0, 6, _sanitize_for_pdf(stripped),
            new_x=XPos.LMARGIN, new_y=YPos.NEXT, markdown=True,
        )

    pdf.output(output_path)
    return output_path


def _add_runs_with_inline_bold(paragraph, text: str) -> None:
    """Splits on **bold** markers and adds each segment as its own run."""
    for i, segment in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not segment:
            continue
        run = paragraph.add_run(segment)
        run.bold = bool(i % 2)  # odd-indexed segments were inside **...**


def markdown_to_docx(markdown_text: str, output_path: str) -> str:
    from docx import Document

    document = Document()

    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()

        if not line:
            continue

        level = _heading_level(line)
        if level:
            text = line.lstrip("#").strip().replace("**", "")
            document.add_heading(text, level=min(level, 9))
            continue

        stripped = line.strip()
        if stripped.startswith(("- ", "* ")):
            _add_runs_with_inline_bold(
                document.add_paragraph(style="List Bullet"), stripped[2:]
            )
            continue

        _add_runs_with_inline_bold(document.add_paragraph(), stripped)

    document.save(output_path)
    return output_path
